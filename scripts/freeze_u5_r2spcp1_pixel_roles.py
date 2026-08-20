#!/usr/bin/env python3
"""Freeze scene-grouped SPCP pixel roles without reading image payloads."""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.audit_u5_r2spcp0_metadata_source_lock import (
    ARCHIVE_URL,
    CENTRAL_OFFSET,
    CENTRAL_SHA256,
    CENTRAL_SIZE,
    canonical_sha256,
    parse_central_directory,
    range_get,
    sha256_bytes,
)

EXPECTED_VARIANTS = (
    "01_01",
    "02_01",
    "02_02",
    "02_03",
    "03_01",
    "03_02",
    "03_03",
    "03_04",
    "04_01",
    "04_02",
    "04_03",
    "04_04",
)
ROLE_COUNTS = (
    ("development", 8),
    ("confirmation", 8),
    ("operator_fit", 128),
    ("operator_calibration", 32),
    ("operator_sealed", 32),
    ("reserve", 249),
)
ORDER_SHA256 = "8c42140ae0f90f37f32706911ab86cca9f377077bbd18ac301262d952bf5f58c"
SCORE_SHA256 = "ee5f0fc830ebd40cab3e25b379aa0e01ec5cc4793a55315504c2d06f6af7900d"


def selection_key(scene_id: str) -> str:
    return hashlib.sha256(f"U5.R2SPCP1\0{scene_id}".encode()).hexdigest()


def eligible_scenes(score_path: Path) -> list[str]:
    if sha256_bytes(score_path.read_bytes()) != SCORE_SHA256:
        raise ValueError("score annotation hash mismatch")
    workbook = load_workbook(score_path, read_only=True, data_only=True)
    rows = workbook.active.iter_rows(values_only=True)
    header = next(rows)
    if header[0] != "image" or len(header) != 21:
        raise ValueError("score workbook header differs")
    variants: dict[str, set[str]] = {}
    for row in rows:
        image_id = row[0][:-4]
        scene_id, family, member = image_id.split("_")
        variants.setdefault(scene_id, set()).add(f"{family}_{member}")
    expected = set(EXPECTED_VARIANTS)
    eligible = [scene for scene, values in variants.items() if values == expected]
    if len(eligible) != 457:
        raise ValueError("single-source eligible scene count mismatch")
    return sorted(eligible, key=lambda scene: (selection_key(scene), scene))


def _verify_order_annotation(path: Path) -> None:
    if sha256_bytes(path.read_bytes()) != ORDER_SHA256:
        raise ValueError("order annotation hash mismatch")


def freeze(order_path: Path, score_path: Path) -> dict[str, Any]:
    _verify_order_annotation(order_path)
    scenes = eligible_scenes(score_path)
    central = range_get(
        ARCHIVE_URL,
        CENTRAL_OFFSET,
        CENTRAL_OFFSET + CENTRAL_SIZE - 1,
    )
    if sha256_bytes(central) != CENTRAL_SHA256:
        raise ValueError("central-directory hash mismatch")
    by_name = {member.name: member for member in parse_central_directory(central)}
    roles: dict[str, Any] = {}
    cursor = 0
    for role, count in ROLE_COUNTS:
        selected = scenes[cursor : cursor + count]
        cursor += count
        image_members = []
        for scene in selected:
            for variant in EXPECTED_VARIANTS:
                name = f"SPCP_dataset/images/{scene}_{variant}.png"
                member = by_name.get(name)
                if member is None:
                    raise ValueError(f"missing canonical image member: {name}")
                image_members.append(
                    {
                        "image_id": f"{scene}_{variant}",
                        "scene_id": scene,
                        "variant_id": variant,
                        "member": name,
                        "crc32": f"{member.crc32:08x}",
                        "method": member.method,
                        "compressed_size": member.compressed_size,
                        "uncompressed_size": member.uncompressed_size,
                        "local_offset": member.local_offset,
                    }
                )
        roles[role] = {
            "scene_count": len(selected),
            "scene_ids": selected,
            "scene_ids_sha256": canonical_sha256(selected),
            "image_member_count": len(image_members),
            "image_members_sha256": canonical_sha256(image_members),
            "image_members": image_members,
        }
    if cursor != len(scenes):
        raise AssertionError("role allocation does not consume exact eligible scene set")
    manifest: dict[str, Any] = {
        "schema": "neuro-film.u5-r2spcp1-pixel-role-lock.v1",
        "status": "LOCKED_BEFORE_IMAGE_MEMBER_PAYLOAD_READ",
        "lock_id": "",
        "dataset_revision": "068af97eed82969f15278db3af4bd450176cf6f3",
        "archive_url": ARCHIVE_URL,
        "central_directory_sha256": CENTRAL_SHA256,
        "order_annotation_sha256": ORDER_SHA256,
        "score_annotation_sha256": SCORE_SHA256,
        "eligible_scene_definition": list(EXPECTED_VARIANTS),
        "eligible_scene_count": len(scenes),
        "eligible_scene_ids_sha256": canonical_sha256(sorted(scenes)),
        "selection_key": "sha256(UTF8('U5.R2SPCP1' + NUL + scene_id)); ascending",
        "roles": roles,
        "image_member_payload_bytes_read": 0,
        "image_pixels_decoded": 0,
        "preference_or_score_values_used": 0,
        "operator_fit_count": 0,
        "render_count": 0,
        "product_capability_opened": False,
    }
    manifest["lock_id"] = canonical_sha256(manifest)
    return manifest


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--order-annotation", type=Path, required=True)
    parser.add_argument("--score-annotation", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.output.exists():
        raise SystemExit("output role lock already exists")
    manifest = freeze(args.order_annotation.resolve(), args.score_annotation.resolve())
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(manifest, indent=2, sort_keys=True, allow_nan=False) + "\n",
        encoding="utf-8",
    )
    print(json.dumps({"lock_id": manifest["lock_id"], "roles": {key: value["scene_count"] for key, value in manifest["roles"].items()}}, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
