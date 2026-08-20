#!/usr/bin/env python3
"""Range-only structural audit for the official SPCP preference dataset."""

from __future__ import annotations

import argparse
import binascii
import hashlib
import json
import math
import re
import struct
import urllib.request
import zlib
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ARCHIVE_URL = (
    "https://huggingface.co/datasets/zwx8981/SPCP_dataset/resolve/"
    "068af97eed82969f15278db3af4bd450176cf6f3/SPCP_dataset.zip"
)
CENTRAL_OFFSET = 9_049_028_973
CENTRAL_SIZE = 1_433_756
CENTRAL_SHA256 = "f3892690a8434fd42a412c4cf952e050d542740d2dea04b4cad976ad2ea1415d"
ARCHIVE_SIZE = 9_050_462_827
EXPECTED_XLSX = {
    "SPCP_dataset/order_trans.xlsx": {
        "sha256": "8c42140ae0f90f37f32706911ab86cca9f377077bbd18ac301262d952bf5f58c",
        "rows": 45_001,
        "columns": 21,
    },
    "SPCP_dataset/score_trans2.xlsx": {
        "sha256": "ee5f0fc830ebd40cab3e25b379aa0e01ec5cc4793a55315504c2d06f6af7900d",
        "rows": 12_001,
        "columns": 21,
    },
}
CANONICAL_PNG = re.compile(r"(?:^|/)(I\d{4}_\d{2}_\d{2})\.png$")
PAIR_ID = re.compile(r"^I\d{4}_\d{2}_\d{2}$")


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def canonical_sha256(value: Any) -> str:
    payload = json.dumps(
        value,
        sort_keys=True,
        separators=(",", ":"),
        ensure_ascii=True,
        allow_nan=False,
    ).encode("utf-8")
    return sha256_bytes(payload)


def range_get(url: str, start: int, end: int) -> bytes:
    if start < 0 or end < start or end >= ARCHIVE_SIZE:
        raise ValueError("invalid bounded archive range")
    request = urllib.request.Request(
        url,
        headers={
            "Range": f"bytes={start}-{end}",
            "User-Agent": "NeuroFilm-U5-R2SPCP0/1.0 (range-only research audit)",
        },
    )
    with urllib.request.urlopen(request, timeout=120) as response:
        status = getattr(response, "status", None)
        content_range = response.headers.get("Content-Range")
        expected_range = f"bytes {start}-{end}/{ARCHIVE_SIZE}"
        if status != 206 or content_range != expected_range:
            raise RuntimeError(
                f"server did not honor exact range: status={status}, "
                f"content_range={content_range!r}"
            )
        expected_length = end - start + 1
        content_length = response.headers.get("Content-Length")
        if content_length is None or int(content_length) != expected_length:
            raise RuntimeError("range Content-Length differs")
        payload = response.read(expected_length + 1)
    if len(payload) != expected_length:
        raise RuntimeError("range payload length differs")
    return payload


@dataclass(frozen=True)
class ZipMember:
    name: str
    flags: int
    method: int
    crc32: int
    compressed_size: int
    uncompressed_size: int
    local_offset: int


def _zip64_values(extra: bytes) -> list[int]:
    cursor = 0
    while cursor + 4 <= len(extra):
        kind, size = struct.unpack_from("<HH", extra, cursor)
        cursor += 4
        value = extra[cursor : cursor + size]
        cursor += size
        if kind == 0x0001:
            if len(value) % 8:
                raise ValueError("invalid ZIP64 extra field")
            return list(struct.unpack(f"<{len(value) // 8}Q", value))
    return []


def parse_central_directory(payload: bytes) -> list[ZipMember]:
    members: list[ZipMember] = []
    cursor = 0
    while cursor < len(payload):
        if payload[cursor : cursor + 4] != b"PK\x01\x02":
            raise ValueError(f"invalid central-directory signature at {cursor}")
        if cursor + 46 > len(payload):
            raise ValueError("truncated central-directory entry")
        values = struct.unpack_from("<4s6H3I5H2I", payload, cursor)
        (
            _,
            _,
            _,
            flags,
            method,
            _,
            _,
            crc32,
            compressed_size,
            uncompressed_size,
            name_length,
            extra_length,
            comment_length,
            _,
            _,
            _,
            local_offset,
        ) = values
        body = cursor + 46
        name_bytes = payload[body : body + name_length]
        extra = payload[body + name_length : body + name_length + extra_length]
        end = body + name_length + extra_length + comment_length
        if end > len(payload):
            raise ValueError("truncated central-directory body")
        name = name_bytes.decode("utf-8" if flags & 0x800 else "cp437")
        zip64 = iter(_zip64_values(extra))
        if uncompressed_size == 0xFFFFFFFF:
            uncompressed_size = next(zip64)
        if compressed_size == 0xFFFFFFFF:
            compressed_size = next(zip64)
        if local_offset == 0xFFFFFFFF:
            local_offset = next(zip64)
        members.append(
            ZipMember(
                name=name,
                flags=flags,
                method=method,
                crc32=crc32,
                compressed_size=compressed_size,
                uncompressed_size=uncompressed_size,
                local_offset=local_offset,
            )
        )
        cursor = end
    return members


def extract_member(url: str, member: ZipMember) -> bytes:
    header = range_get(url, member.local_offset, member.local_offset + 29)
    values = struct.unpack("<4s5H3I2H", header)
    if values[0] != b"PK\x03\x04":
        raise ValueError(f"invalid local header: {member.name}")
    flags, method = values[2], values[3]
    name_length, extra_length = values[-2], values[-1]
    if flags != member.flags or method != member.method:
        raise ValueError(f"local/central metadata mismatch: {member.name}")
    variable = range_get(
        url,
        member.local_offset + 30,
        member.local_offset + 30 + name_length + extra_length - 1,
    )
    name = variable[:name_length].decode("utf-8" if flags & 0x800 else "cp437")
    if name != member.name:
        raise ValueError(f"local/central filename mismatch: {member.name}")
    data_start = member.local_offset + 30 + name_length + extra_length
    compressed = range_get(
        url,
        data_start,
        data_start + member.compressed_size - 1,
    )
    if method == 0:
        raw = compressed
    elif method == 8:
        raw = zlib.decompress(compressed, -zlib.MAX_WBITS)
    else:
        raise ValueError(f"unsupported ZIP method {method}: {member.name}")
    if len(raw) != member.uncompressed_size:
        raise ValueError(f"uncompressed size mismatch: {member.name}")
    if binascii.crc32(raw) & 0xFFFFFFFF != member.crc32:
        raise ValueError(f"CRC mismatch: {member.name}")
    return raw


def workbook_facts(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if len(workbook.sheetnames) != 1:
        raise ValueError(f"workbook sheet count differs: {path.name}")
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"empty workbook: {path.name}")
    return {
        "sheet_name": sheet.title,
        "row_count": len(rows),
        "column_count": max(len(row) for row in rows),
        "header": list(rows[0]),
        "first_data_rows": [list(row) for row in rows[1:4]],
        "all_cells_nonempty": all(value is not None for row in rows for value in row),
    }


def _connected(nodes: set[str], edges: set[tuple[str, str]]) -> bool:
    if not nodes:
        return False
    adjacency = {node: set() for node in nodes}
    for left, right in edges:
        adjacency[left].add(right)
        adjacency[right].add(left)
    pending = [min(nodes)]
    visited: set[str] = set()
    while pending:
        node = pending.pop()
        if node in visited:
            continue
        visited.add(node)
        pending.extend(sorted(adjacency[node] - visited))
    return visited == nodes


def annotation_graph_facts(order_path: Path, score_path: Path) -> dict[str, Any]:
    score_book = load_workbook(score_path, read_only=True, data_only=True)
    score_sheet = score_book[score_book.sheetnames[0]]
    score_rows = score_sheet.iter_rows(values_only=True)
    score_header = next(score_rows)
    if score_header[0] != "image" or len(score_header) != 21:
        raise ValueError("score annotation header differs")
    score_ids: set[str] = set()
    score_scenes: dict[str, set[str]] = {}
    score_min = math.inf
    score_max = -math.inf
    for row in score_rows:
        if len(row) != 21 or not isinstance(row[0], str) or not row[0].endswith(".png"):
            raise ValueError("invalid score annotation row")
        image_id = row[0][:-4]
        if not PAIR_ID.fullmatch(image_id) or image_id in score_ids:
            raise ValueError("invalid or duplicate score image identity")
        values = row[1:]
        if any(not isinstance(value, (int, float)) or not math.isfinite(value) for value in values):
            raise ValueError("non-finite score annotation")
        score_ids.add(image_id)
        score_scenes.setdefault(image_id.split("_")[0], set()).add(image_id)
        score_min = min(score_min, *(float(value) for value in values))
        score_max = max(score_max, *(float(value) for value in values))

    order_book = load_workbook(order_path, read_only=True, data_only=True)
    order_sheet = order_book[order_book.sheetnames[0]]
    order_rows = order_sheet.iter_rows(values_only=True)
    order_header = next(order_rows)
    if order_header[0] != "pair" or len(order_header) != 21:
        raise ValueError("pair annotation header differs")
    pair_scenes: dict[str, set[tuple[str, str]]] = {}
    endpoint_scenes: dict[str, set[str]] = {}
    label_counts = {0: 0, 1: 0}
    majority_tie_count = 0
    for row in order_rows:
        if len(row) != 21 or not isinstance(row[0], str):
            raise ValueError("invalid pair annotation row")
        endpoints = row[0].split(",")
        if len(endpoints) != 2 or any(not PAIR_ID.fullmatch(value) for value in endpoints):
            raise ValueError("invalid pair endpoint identity")
        left, right = endpoints
        left_scene, right_scene = left.split("_")[0], right.split("_")[0]
        if left_scene != right_scene or left == right:
            raise ValueError("cross-scene or self pair")
        edge = tuple(sorted((left, right)))
        edges = pair_scenes.setdefault(left_scene, set())
        if edge in edges:
            raise ValueError("duplicate unordered pair edge")
        edges.add(edge)
        endpoint_scenes.setdefault(left_scene, set()).update((left, right))
        labels = row[1:]
        if any(type(value) is not int or value not in (0, 1) for value in labels):
            raise ValueError("pair preference label outside exact integer domain")
        ones = sum(labels)
        label_counts[0] += 20 - ones
        label_counts[1] += ones
        majority_tie_count += int(ones == 10)

    scene_ids = sorted(score_scenes)
    if set(pair_scenes) != set(scene_ids) or set(endpoint_scenes) != set(scene_ids):
        raise ValueError("pair and score scene sets differ")
    pair_counts = [len(pair_scenes[scene]) for scene in scene_ids]
    node_counts = [len(endpoint_scenes[scene]) for scene in scene_ids]
    score_counts = [len(score_scenes[scene]) for scene in scene_ids]
    connected = [_connected(endpoint_scenes[scene], pair_scenes[scene]) for scene in scene_ids]
    endpoint_matches = [endpoint_scenes[scene] == score_scenes[scene] for scene in scene_ids]
    return {
        "scene_count": len(scene_ids),
        "score_image_count": len(score_ids),
        "score_images_per_scene_min": min(score_counts),
        "score_images_per_scene_max": max(score_counts),
        "score_value_min": score_min,
        "score_value_max": score_max,
        "pair_row_count": sum(pair_counts),
        "pair_rows_per_scene_min": min(pair_counts),
        "pair_rows_per_scene_max": max(pair_counts),
        "pair_nodes_per_scene_min": min(node_counts),
        "pair_nodes_per_scene_max": max(node_counts),
        "connected_scene_count": sum(connected),
        "endpoint_set_matches_score_set_scene_count": sum(endpoint_matches),
        "cross_scene_pair_count": 0,
        "preference_label_counts": label_counts,
        "majority_tie_edge_count": majority_tie_count,
        "complete_graph_edges_for_12_nodes": 66,
        "released_edges_per_scene": min(pair_counts),
        "released_graph_is_complete": all(value == 66 for value in pair_counts),
        "released_pair_density": min(pair_counts) / 66.0,
        "scene_ids_sha256": canonical_sha256(scene_ids),
        "score_image_ids_sha256": canonical_sha256(sorted(score_ids)),
        "pair_edges_sha256": canonical_sha256(
            [
                [scene, left, right]
                for scene in scene_ids
                for left, right in sorted(pair_scenes[scene])
            ]
        ),
    }


def run(output_root: Path) -> dict[str, Any]:
    output_root.mkdir(parents=True, exist_ok=False)
    central = range_get(
        ARCHIVE_URL,
        CENTRAL_OFFSET,
        CENTRAL_OFFSET + CENTRAL_SIZE - 1,
    )
    if sha256_bytes(central) != CENTRAL_SHA256:
        raise ValueError("central-directory hash mismatch")
    members = parse_central_directory(central)
    by_name = {member.name: member for member in members}
    if len(by_name) != len(members):
        raise ValueError("duplicate ZIP member names")

    annotation_facts: dict[str, Any] = {}
    for filename, expected in EXPECTED_XLSX.items():
        member = by_name.get(filename)
        if member is None:
            raise ValueError(f"missing annotation member: {filename}")
        raw = extract_member(ARCHIVE_URL, member)
        if sha256_bytes(raw) != expected["sha256"]:
            raise ValueError(f"annotation SHA mismatch: {filename}")
        path = output_root / Path(filename).name
        path.write_bytes(raw)
        facts = workbook_facts(path)
        if facts["row_count"] != expected["rows"]:
            raise ValueError(f"annotation row count mismatch: {filename}")
        if facts["column_count"] != expected["columns"]:
            raise ValueError(f"annotation column count mismatch: {filename}")
        annotation_facts[filename] = {
            "member": member.__dict__,
            "sha256": sha256_bytes(raw),
            **facts,
        }

    png_names = [member.name for member in members if member.name.lower().endswith(".png")]
    canonical_ids = []
    noncanonical_png = []
    for name in png_names:
        match = CANONICAL_PNG.search(name)
        if match:
            canonical_ids.append(match.group(1))
        else:
            noncanonical_png.append(name)
    graph = annotation_graph_facts(
        output_root / "order_trans.xlsx",
        output_root / "score_trans2.xlsx",
    )
    canonical_set = set(canonical_ids)
    duplicate_facts = []
    for name in sorted(noncanonical_png):
        canonical_name = name.replace("(1).png", ".png")
        duplicate = by_name[name]
        canonical = by_name.get(canonical_name)
        if canonical is None:
            raise ValueError(f"noncanonical PNG has no canonical peer: {name}")
        duplicate_facts.append(
            {
                "member": name,
                "canonical_member": canonical_name,
                "same_crc32": duplicate.crc32 == canonical.crc32,
                "same_compressed_size": duplicate.compressed_size == canonical.compressed_size,
                "same_uncompressed_size": duplicate.uncompressed_size == canonical.uncompressed_size,
            }
        )
    gates = {
        "scene_count_exact": graph["scene_count"] == 1000,
        "variants_per_scene_exact": (
            graph["score_images_per_scene_min"] == graph["score_images_per_scene_max"] == 12
        ),
        "canonical_image_count_exact": len(canonical_set) == 12000,
        "pair_rows_per_scene_exact": (
            graph["pair_rows_per_scene_min"] == graph["pair_rows_per_scene_max"] == 45
        ),
        "pair_row_count_exact": graph["pair_row_count"] == 45000,
        "subjects_per_row_exact": all(
            facts["column_count"] == 21 for facts in annotation_facts.values()
        ),
        "preference_label_domain_exact": set(graph["preference_label_counts"]) == {0, 1},
        "cross_scene_pair_count_exact": graph["cross_scene_pair_count"] == 0,
        "each_scene_pair_graph_connected": graph["connected_scene_count"] == 1000,
        "pair_endpoints_equal_scored_image_set": (
            graph["endpoint_set_matches_score_set_scene_count"] == 1000
        ),
        "canonical_png_equals_scored_image_set": (
            graph["score_image_count"] == len(canonical_set)
            and graph["score_image_ids_sha256"]
            == canonical_sha256(sorted(canonical_set))
        ),
        "duplicate_or_noncanonical_png_members_reported_and_excluded": (
            len(duplicate_facts) == 3
            and all(
                row["same_crc32"]
                and row["same_compressed_size"]
                and row["same_uncompressed_size"]
                for row in duplicate_facts
            )
        ),
        "image_payload_bytes_read_exact": True,
        "operator_fit_count_exact": True,
        "render_count_exact": True,
        "scientific_score_count_exact": True,
    }
    report: dict[str, Any] = {
        "schema": "neuro-film.u5-r2spcp0-range-audit-result.v1",
        "status": "PASS_METADATA_ONLY_SOURCE_AND_ANNOTATION_ELIGIBILITY" if all(gates.values()) else "FAIL_CLOSED_METADATA_SOURCE_GATE",
        "archive_url": ARCHIVE_URL,
        "archive_size_bytes": ARCHIVE_SIZE,
        "central_directory": {
            "offset": CENTRAL_OFFSET,
            "size": CENTRAL_SIZE,
            "sha256": sha256_bytes(central),
            "member_count": len(members),
        },
        "inventory": {
            "png_member_count": len(png_names),
            "canonical_png_identity_count": len(set(canonical_ids)),
            "canonical_png_member_count": len(canonical_ids),
            "noncanonical_png_members": sorted(noncanonical_png),
            "noncanonical_duplicate_facts": duplicate_facts,
        },
        "annotations": annotation_facts,
        "graph": graph,
        "publication_release_scope_note": "The release has 45 connected edges over 12 nodes per scene, not the 66 edges of a complete 12-node graph; future work must call it a connected sampled pair graph rather than a complete exhaustive graph.",
        "gates": gates,
        "reads": {
            "image_member_payload_bytes": 0,
            "operator_fit_count": 0,
            "render_count": 0,
            "scientific_score_count": 0,
        },
        "claim_ceiling": "Structurally complete connected same-scene direct human colour-preference annotations for a later private explicit-operator D0; no image rights expansion, operator result, model, product or commercial claim.",
    }
    report["report_id"] = canonical_sha256(report)
    (output_root / "inspection.json").write_text(
        json.dumps(report, indent=2, sort_keys=True, allow_nan=False) + "\n",
        encoding="utf-8",
    )
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-root", type=Path, required=True)
    args = parser.parse_args()
    report = run(args.output_root.resolve())
    print(
        json.dumps(
            {
                "report_id": report["report_id"],
                "member_count": report["central_directory"]["member_count"],
                "inventory": report["inventory"],
            },
            sort_keys=True,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
