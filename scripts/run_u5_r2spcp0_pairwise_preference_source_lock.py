#!/usr/bin/env python3
"""Range-only SPCP preference-source integrity and structure audit."""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import re
import struct
import urllib.request
import zlib
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CONTRACT = ROOT / "configs/u5_r2spcp0_pairwise_preference_source_lock_v1.json"
DEFAULT_OUTPUT = (
    ROOT
    / "outputs/eval/u5_r2spcp0_pairwise_preference_source_lock_v1/formal_report.json"
)


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _canonical_bytes(payload: Any) -> bytes:
    return (json.dumps(payload, sort_keys=True, separators=(",", ":")) + "\n").encode()


def _fetch(url: str, byte_range: tuple[int, int] | None = None) -> tuple[bytes, dict[str, str]]:
    headers = {"User-Agent": "neuro-film-u5-r2spcp0/1"}
    if byte_range is not None:
        headers["Range"] = f"bytes={byte_range[0]}-{byte_range[1]}"
    request = urllib.request.Request(url, headers=headers)
    with urllib.request.urlopen(request, timeout=120) as response:
        data = response.read()
        response_headers = {key.lower(): value for key, value in response.headers.items()}
        status = int(response.status)
    if byte_range is not None:
        expected = byte_range[1] - byte_range[0] + 1
        if status != 206 or len(data) != expected:
            raise RuntimeError(
                f"range response drift: status={status}, bytes={len(data)}, expected={expected}"
            )
        expected_range = f"bytes {byte_range[0]}-{byte_range[1]}/"
        if not response_headers.get("content-range", "").startswith(expected_range):
            raise RuntimeError("content-range drift")
    return data, response_headers


def _parse_central_directory(data: bytes) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    offset = 0
    while offset < len(data):
        if data[offset : offset + 4] != b"PK\x01\x02":
            raise ValueError(f"central directory signature drift at {offset}")
        values = struct.unpack_from("<IHHHHHHIIIHHHHHII", data, offset)
        (
            _,
            _,
            _,
            _,
            method,
            _,
            _,
            crc32,
            compressed_size,
            uncompressed_size,
            name_len,
            extra_len,
            comment_len,
            _,
            _,
            _,
            local_offset,
        ) = values
        name_start = offset + 46
        name = data[name_start : name_start + name_len].decode("utf-8")
        extra_start = name_start + name_len
        extra = data[extra_start : extra_start + extra_len]
        if (
            uncompressed_size == 0xFFFFFFFF
            or compressed_size == 0xFFFFFFFF
            or local_offset == 0xFFFFFFFF
        ):
            cursor = 0
            zip64_payload: bytes | None = None
            while cursor < len(extra):
                field_id, field_size = struct.unpack_from("<HH", extra, cursor)
                field = extra[cursor + 4 : cursor + 4 + field_size]
                if field_id == 0x0001:
                    zip64_payload = field
                    break
                cursor += 4 + field_size
            if zip64_payload is None:
                raise ValueError(f"ZIP64 extra missing for {name}")
            cursor = 0
            if uncompressed_size == 0xFFFFFFFF:
                uncompressed_size = struct.unpack_from("<Q", zip64_payload, cursor)[0]
                cursor += 8
            if compressed_size == 0xFFFFFFFF:
                compressed_size = struct.unpack_from("<Q", zip64_payload, cursor)[0]
                cursor += 8
            if local_offset == 0xFFFFFFFF:
                local_offset = struct.unpack_from("<Q", zip64_payload, cursor)[0]
        rows.append(
            {
                "name": name,
                "method": method,
                "crc32": crc32,
                "compressed_size": compressed_size,
                "uncompressed_size": uncompressed_size,
                "local_offset": local_offset,
            }
        )
        offset = extra_start + extra_len + comment_len
    return rows


def _extract_member(
    range_bytes: bytes,
    *,
    range_start: int,
    row: dict[str, Any],
) -> bytes:
    local = int(row["local_offset"]) - range_start
    if range_bytes[local : local + 4] != b"PK\x03\x04":
        raise ValueError(f"local header missing for {row['name']}")
    values = struct.unpack_from("<IHHHHHIIIHH", range_bytes, local)
    method, name_len, extra_len = values[3], values[9], values[10]
    if method != 8:
        raise ValueError(f"unsupported compression method for {row['name']}: {method}")
    payload_start = local + 30 + name_len + extra_len
    payload_end = payload_start + int(row["compressed_size"])
    compressed = range_bytes[payload_start:payload_end]
    if len(compressed) != int(row["compressed_size"]):
        raise ValueError(f"compressed member truncated: {row['name']}")
    decoded = zlib.decompress(compressed, -15)
    if len(decoded) != int(row["uncompressed_size"]):
        raise ValueError(f"uncompressed member size drift: {row['name']}")
    if zlib.crc32(decoded) & 0xFFFFFFFF != int(row["crc32"]):
        raise ValueError(f"member CRC drift: {row['name']}")
    return decoded


def _sheet_rows(data: bytes) -> tuple[tuple[Any, ...], list[tuple[Any, ...]]]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    if workbook.sheetnames != ["Sheet1"]:
        raise ValueError("workbook sheet inventory drift")
    worksheet = workbook["Sheet1"]
    iterator = worksheet.iter_rows(values_only=True)
    header = tuple(next(iterator))
    rows = [tuple(row) for row in iterator]
    workbook.close()
    return header, rows


def _graph_connected(nodes: set[str], edges: list[tuple[str, str]]) -> bool:
    if not nodes:
        return False
    adjacency: dict[str, set[str]] = defaultdict(set)
    for left, right in edges:
        adjacency[left].add(right)
        adjacency[right].add(left)
    seen: set[str] = set()
    stack = [next(iter(nodes))]
    while stack:
        node = stack.pop()
        if node in seen:
            continue
        seen.add(node)
        stack.extend(adjacency[node] - seen)
    return seen == nodes


def _stable_id(payload: dict[str, Any]) -> str:
    stable = dict(payload)
    stable.pop("stable_evidence_id", None)
    return f"sha256:{_sha256(_canonical_bytes(stable))}"


def _report_schema(experiment_id: str) -> str:
    token = experiment_id.lower().replace(".", "-")
    return f"neuro-film.{token}-pairwise-preference-source-lock-report.v1"


def run(contract_path: Path, output_path: Path) -> dict[str, Any]:
    contract_bytes = contract_path.read_bytes()
    contract = json.loads(contract_bytes)
    source = contract["source"]
    ranges = contract["range_protocol"]
    gates = contract["frozen_gates"]
    parent = contract.get("parent")
    parent_facts: dict[str, Any] | None = None
    if parent is not None:
        decision_path = ROOT / parent["decision_path"]
        evidence_path = ROOT / parent["evidence_path"]
        decision_bytes = decision_path.read_bytes()
        evidence_bytes = evidence_path.read_bytes()
        decision_payload = json.loads(decision_bytes)
        parent_facts = {
            "decision_sha256": _sha256(decision_bytes),
            "evidence_sha256": _sha256(evidence_bytes),
            "decision": decision_payload.get("decision"),
        }

    api_bytes, _ = _fetch(source["api_tree_url"])
    readme_bytes, _ = _fetch(source["readme_url"])
    tail, _ = _fetch(
        source["zip_url"], (ranges["tail_start"], ranges["tail_end_inclusive"])
    )
    annotations, _ = _fetch(
        source["zip_url"],
        (ranges["annotation_start"], ranges["annotation_end_inclusive"]),
    )

    api_tree = json.loads(api_bytes)
    central_size = int(ranges["expected_central_directory_size"])
    central = tail[:central_size]
    members = _parse_central_directory(central)
    member_by_name = {row["name"]: row for row in members}

    pair_member = next(
        row for row in members if row["name"].endswith(contract["annotations"]["pairwise"]["member"])
    )
    score_member = next(
        row for row in members if row["name"].endswith(contract["annotations"]["scores"]["member"])
    )
    pair_bytes = _extract_member(
        annotations,
        range_start=int(ranges["annotation_start"]),
        row=pair_member,
    )
    score_bytes = _extract_member(
        annotations,
        range_start=int(ranges["annotation_start"]),
        row=score_member,
    )

    pair_header, pair_rows = _sheet_rows(pair_bytes)
    _score_header, score_rows = _sheet_rows(score_bytes)
    subject_columns = len(pair_header) - 1

    canonical_re = re.compile(r"^SPCP_dataset/images/(I\d{4}_\d{2}_\d{2})\.png$")
    canonical_members = [row for row in members if canonical_re.fullmatch(row["name"])]
    png_members = [row for row in members if row["name"].lower().endswith(".png")]
    extra_members = [row for row in png_members if not canonical_re.fullmatch(row["name"])]
    canonical_names = {
        canonical_re.fullmatch(row["name"]).group(1) + ".png"
        for row in canonical_members
    }
    scene_counts = Counter(name.split("_")[0] for name in canonical_names)

    score_names = [str(row[0]) for row in score_rows]
    score_values = [value for row in score_rows for value in row[1:]]
    pair_ids = [str(row[0]) for row in pair_rows]
    pair_values = [value for row in pair_rows for value in row[1:]]
    edges_by_scene: dict[str, list[tuple[str, str]]] = defaultdict(list)
    cross_scene_pairs = 0
    pair_endpoint_missing = 0
    for pair_id in pair_ids:
        left, right = pair_id.split(",")
        left_scene, right_scene = left.split("_")[0], right.split("_")[0]
        if left_scene != right_scene:
            cross_scene_pairs += 1
        if left + ".png" not in canonical_names or right + ".png" not in canonical_names:
            pair_endpoint_missing += 1
        edges_by_scene[left_scene].append((left, right))
    connected_graphs = sum(
        _graph_connected(
            {name.removesuffix(".png") for name in canonical_names if name.startswith(scene + "_")},
            edges,
        )
        for scene, edges in edges_by_scene.items()
    )

    extras_duplicate_exact = 0
    extra_names: list[str] = []
    for row in extra_members:
        extra_names.append(row["name"])
        base_name = row["name"].replace("(1).png", ".png")
        base = member_by_name.get(base_name)
        if base and all(
            row[key] == base[key]
            for key in ("crc32", "compressed_size", "uncompressed_size")
        ):
            extras_duplicate_exact += 1

    pair_expected = contract["annotations"]["pairwise"]
    score_expected = contract["annotations"]["scores"]
    facts = {
        "api_tree_sha256": _sha256(api_bytes),
        "root_files": sorted(
            [
                {"path": row["path"], "size": row["size"], "oid": row["oid"]}
                for row in api_tree
            ],
            key=lambda row: row["path"],
        ),
        "readme_bytes": len(readme_bytes),
        "readme_sha256": _sha256(readme_bytes),
        "readme_exact": readme_bytes.decode("utf-8"),
        "tail_bytes_read": len(tail),
        "tail_sha256": _sha256(tail),
        "central_directory_bytes": len(central),
        "central_directory_sha256": _sha256(central),
        "zip_member_count": len(members),
        "canonical_image_count": len(canonical_members),
        "scene_count": len(scene_counts),
        "variants_per_scene_min": min(scene_counts.values()),
        "variants_per_scene_max": max(scene_counts.values()),
        "duplicate_extra_count": len(extra_members),
        "duplicate_extras_match_canonical_crc_and_sizes": extras_duplicate_exact,
        "duplicate_extra_names": sorted(extra_names),
        "pair_annotation_sha256": _sha256(pair_bytes),
        "score_annotation_sha256": _sha256(score_bytes),
        "pair_rows": len(pair_rows),
        "score_rows": len(score_rows),
        "subject_columns": subject_columns,
        "pair_ids_unique": len(set(pair_ids)) == len(pair_ids),
        "score_ids_unique": len(set(score_names)) == len(score_names),
        "score_ids_equal_canonical_members": set(score_names) == canonical_names,
        "pair_label_domain": sorted(set(pair_values)),
        "pair_missing_values": sum(value is None for value in pair_values),
        "score_nonfinite_values": sum(
            value is None or not math.isfinite(float(value)) for value in score_values
        ),
        "cross_scene_pairs": cross_scene_pairs,
        "pair_endpoint_missing": pair_endpoint_missing,
        "connected_pair_graphs": connected_graphs,
        "pairs_per_scene_min": min(map(len, edges_by_scene.values())),
        "pairs_per_scene_max": max(map(len, edges_by_scene.values())),
        "zip_bytes_read_total": len(tail) + len(annotations),
        "image_payload_bytes_read": 0,
    }
    if parent_facts is not None:
        facts["parent"] = parent_facts
    gate_results = {
        "root_file_count_exact": len(api_tree) == gates["root_file_count_exact"],
        "readme_exact": readme_bytes.decode("utf-8")
        == contract["rights"]["required_card_front_matter"],
        "zip_member_count_exact": len(members) == gates["zip_member_count_exact"],
        "central_directory_sha256_exact": facts["central_directory_sha256"]
        == source.get("central_directory_sha256", facts["central_directory_sha256"]),
        "canonical_image_count_exact": len(canonical_members)
        == gates["canonical_image_count_exact"],
        "scene_count_exact": len(scene_counts) == gates["scene_count_exact"],
        "variants_per_scene_exact": set(scene_counts.values())
        == {gates["variants_per_scene_exact"]},
        "duplicate_extra_count_exact": len(extra_members)
        == gates["duplicate_extra_count_exact"],
        "duplicate_extras_match_canonical": extras_duplicate_exact == len(extra_members),
        "pair_annotation_size_crc_exact": all(
            [
                pair_member["compressed_size"] == pair_expected["compressed_size"],
                pair_member["uncompressed_size"] == pair_expected["uncompressed_size"],
                f"{pair_member['crc32']:08x}" == pair_expected["crc32_hex"],
            ]
        ),
        "score_annotation_size_crc_exact": all(
            [
                score_member["compressed_size"] == score_expected["compressed_size"],
                score_member["uncompressed_size"] == score_expected["uncompressed_size"],
                f"{score_member['crc32']:08x}" == score_expected["crc32_hex"],
            ]
        ),
        "pair_annotation_sha256_exact": facts["pair_annotation_sha256"]
        == pair_expected["sha256"],
        "score_annotation_sha256_exact": facts["score_annotation_sha256"]
        == score_expected["sha256"],
        "annotation_shape_exact": len(pair_rows) == gates["pair_rows_exact"]
        and len(score_rows) == gates["score_rows_exact"]
        and subject_columns == gates["subject_columns_exact"],
        "annotation_values_valid": set(pair_values) == set(gates["pair_label_domain"])
        and facts["pair_missing_values"] == 0
        and facts["score_nonfinite_values"] == 0,
        "pair_graph_exact": cross_scene_pairs == gates["cross_scene_pairs_exact"]
        and pair_endpoint_missing == 0
        and connected_graphs == gates["connected_pair_graphs_exact"]
        and facts["pairs_per_scene_min"] == gates["pairs_per_scene_exact"]
        and facts["pairs_per_scene_max"] == gates["pairs_per_scene_exact"],
        "score_endpoint_set_exact": facts["score_ids_equal_canonical_members"],
        "range_budget_exact": facts["zip_bytes_read_total"]
        <= ranges["maximum_total_zip_bytes_read"],
        "image_payload_bytes_read_exact": facts["image_payload_bytes_read"]
        == gates["image_payload_bytes_read_exact"],
    }
    if parent is not None:
        gate_results["parent_exact"] = bool(
            parent_facts
            and parent_facts["decision_sha256"] == parent["decision_sha256"]
            and parent_facts["evidence_sha256"] == parent["evidence_sha256"]
            and parent_facts["decision"] == parent["required_decision"]
        )
    failed = sorted(name for name, passed in gate_results.items() if not passed)
    report = {
        "schema": _report_schema(contract["experiment_id"]),
        "experiment_id": contract["experiment_id"],
        "contract_path": contract_path.relative_to(ROOT).as_posix(),
        "contract_sha256": _sha256(contract_bytes),
        "facts": facts,
        "gates": gate_results,
        "failed_gates": failed,
        "status": (
            "PASS_METADATA_ONLY_SOURCE_LOCK"
            if not failed
            else "FAIL_CLOSED_METADATA_ONLY_SOURCE_LOCK"
        ),
        "decision": (
            "open_separately_preregistered_explicit_operator_d0"
            if not failed
            else "close_exact_spcp0_protocol_before_image_acquisition"
        ),
        "claim_ceiling": contract["claim_ceiling"],
    }
    report["stable_evidence_id"] = _stable_id(report)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_bytes(_canonical_bytes(report))
    return report


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--contract", type=Path, default=DEFAULT_CONTRACT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    report = run(args.contract.resolve(), args.output.resolve())
    print(json.dumps(report, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
